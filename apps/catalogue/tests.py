from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import Workbook

from apps.catalogue.datafile import metal_family, parse_rows, apply_rows as apply_datafile
from apps.catalogue.intake import infer_metal_name, parse_jewellery_workbook
from apps.catalogue.models import strip_dflt_prefix
from apps.locations.models import Location, LocationType


class IntakeParseTests(TestCase):
    def test_dflt_prefix_is_not_a_metal(self):
        self.assertEqual(strip_dflt_prefix("DFLT - 18K"), "18K")
        self.assertEqual(infer_metal_name("", "DFLT - 18K"), "Gold")
        self.assertEqual(infer_metal_name("", "DFLT - PT900"), "Platinum")
        self.assertEqual(infer_metal_name("", ""), "")

    def test_catches_jewellery_excel_header_on_row_2(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Jewellery Excel"
        ws.append(["title"] * 5)
        ws.append(["supplier", "Supplier_product_code", "Supplier_nick_name", "Metal Purity", "PJNUMBER"])
        ws.append(["YZC", "Z-R4442", "YZC1", "DFLT - 18K", "PJ24881"])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        parsed = parse_jewellery_workbook(buf)
        self.assertEqual(parsed.errors, [])
        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.rows[0].pj, "PJ24881")
        self.assertEqual(parsed.rows[0].purity, "DFLT - 18K")
        self.assertEqual(parsed.rows[0].reference, "Z-R4442")


class IntakeApplyTests(TestCase):
    def test_blank_metal_name_stores_gold_not_silver(self):
        from apps.catalogue.intake import IntakeRow, apply_rows

        Location.objects.create(name="Head office", code="HO", location_type=LocationType.HEAD_OFFICE)
        row = IntakeRow(
            row_number=3,
            pj="PJ24881",
            reference="Z-R4442",
            owner="YZC1",
            supplier="YZC",
            subcategory="RN",
            metal_name="",
            purity="DFLT - 18K",
            weight="1.58",
            actual_price=Decimal("239.4"),
            currency="USD",
            rate=Decimal("63"),
            payment_type="markuptype1",
            markup=Decimal("10"),
            markup_amount=None,
            size="",
            diamond_weight="",
        )
        report = apply_rows([row], Location.objects.get(code="HO"))
        self.assertEqual(report["created"], 1)
        item = Location.objects.get(code="HO").items.get(barcode="PJ24881")
        self.assertEqual(item.product.metal.name, "Gold")
        self.assertEqual(item.product.purity.name, "18K")
        self.assertEqual(item.product.gold_weight, "1.58")
        self.assertEqual(item.product.selling_price, Decimal("215.46"))
        self.assertEqual(len(report["lines"]), 1)
        self.assertEqual(report["lines"][0]["action"], "created")


class IntakeTemplateTests(TestCase):
    def test_sample_workbook_is_the_jewellery_upload(self):
        from io import BytesIO

        from apps.catalogue.intake import TEMPLATE_HEADERS, build_jewellery_template, parse_jewellery_workbook

        parsed = parse_jewellery_workbook(BytesIO(build_jewellery_template()))
        self.assertEqual(parsed.sheet, "Jewellery Excel")
        self.assertEqual(parsed.rows, [])
        self.assertEqual(parsed.errors, ["No data rows under the header."])

        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(build_jewellery_template()))
        headers = [cell.value for cell in wb["Jewellery Excel"][2]]
        self.assertEqual(tuple(headers), TEMPLATE_HEADERS)
        self.assertIn("PJNUMBER", headers)
        self.assertIn("Supplier_product_code", headers)


class IntakeHistoryTests(TestCase):
    def setUp(self):
        from apps.accounts.models import Employee
        from apps.locations.models import Location, LocationType

        self.loc = Location.objects.create(name="Head office", code="HO", location_type=LocationType.HEAD_OFFICE)
        self.user = Employee.objects.create_user(
            employee_code="1001", password="test-pass", first_name="Test"
        )
        self.client.force_login(self.user)

    def _workbook(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Jewellery Excel"
        ws.append(["title"] * 8)
        ws.append([
            "supplier", "Supplier_product_code", "Supplier_nick_name",
            "Metal Purity", "metal_weight", "purchase_type", "due_date", "PJNUMBER",
        ])
        ws.append(["YZC", "Z-R4442", "YZC1", "DFLT - 18K", "1.58", "consignment", "11-11-2026", "PJ24881"])
        buf = BytesIO()
        wb.save(buf)
        return SimpleUploadedFile(
            "YZC.xlsx",
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_upload_lands_on_history_and_keeps_the_file(self):
        from datetime import date

        from django.test import override_settings

        from apps.catalogue.models import ProductIntakeBatch
        from apps.inventory.models import ProductItem

        static = {
            "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        }
        with override_settings(DEBUG=True, STORAGES=static):
            page = self.client.get("/products/add/")
            self.assertEqual(page.status_code, 200)
            self.assertContains(page, "Download sample Excel")
            self.assertContains(page, "Upload history")

            sample = self.client.get("/products/add/template/")
            self.assertEqual(sample.status_code, 200)
            self.assertIn(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                sample["Content-Type"],
            )
            self.assertIn("JewelleryExcelDData.xlsx", sample["Content-Disposition"])

            posted = self.client.post(
                "/products/add/",
                {"location": self.loc.pk, "workbook": self._workbook()},
            )
            self.assertEqual(posted.status_code, 302)
            batch = ProductIntakeBatch.objects.get()
            self.assertEqual(batch.serial_no, 1)
            self.assertEqual(batch.created_count, 1)
            self.assertEqual(batch.filename, "YZC.xlsx")
            self.assertTrue(batch.workbook)
            self.assertEqual(posted["Location"], f"/products/add/history/{batch.pk}/")

            detail = self.client.get(posted["Location"])
            self.assertEqual(detail.status_code, 200)
            self.assertContains(detail, "PJ24881")
            self.assertContains(detail, "Upload 1")
            item = ProductItem.objects.get(barcode="PJ24881")
            self.assertEqual(item.product.product_type, "consignment")
            self.assertEqual(item.product.due_date, date(2026, 11, 11))
            self.assertEqual(item.product.metal.name, "Gold")

            listed = self.client.get("/products/add/")
            self.assertContains(listed, "YZC.xlsx")
            self.assertContains(listed, "View")


class DatafileSyncTests(TestCase):
    def test_blank_metal_type_uses_purity_not_silver(self):
        self.assertEqual(metal_family("", "18K+3G"), "Gold")
        self.assertEqual(metal_family("SL", "925"), "Silver")

    def test_sheet_purity_replaces_wrong_metal(self):
        from apps.catalogue.models import Category, Currency, Metal, ProductMaster
        from apps.inventory.models import ProductItem
        from apps.locations.models import Location, LocationType

        cat = Category.objects.create(name="Jewellery", code="JW")
        cur = Currency.objects.create(code="USD")
        silver = Metal.objects.create(name="Silver")
        loc = Location.objects.create(name="Head office", code="HO", location_type=LocationType.HEAD_OFFICE)
        product = ProductMaster.objects.create(
            name="YZC1", reference_id="BR0704-1-3G", category=cat, currency=cur, metal=silver,
        )
        ProductItem.objects.create(barcode="PJ25073", product=product, location=loc)
        csv_text = (
            "RFID Tag,Metal Type,Metal Purity,Net Weight,Metal Weight,Gross Weight\n"
            "PJ25073,,18K+3G,9.05g,,\n"
        )
        report = apply_datafile(parse_rows(csv_text))
        product.refresh_from_db()
        self.assertEqual(report["updated"], 1)
        self.assertEqual(product.metal.name, "Gold")
        self.assertEqual(product.purity.name, "18K+3G")
        self.assertEqual(product.net_weight, Decimal("9.05"))


class LegacyDateParseTests(TestCase):
    def test_day_month_year_and_excel_serial(self):
        from datetime import date

        from apps.core.legacy_import import as_date, as_purchase_type

        self.assertEqual(as_date("30-08-2026 00:00:00"), date(2026, 8, 30))
        self.assertEqual(as_date("46264"), date(2026, 8, 30))
        self.assertEqual(as_date(46203), date(2026, 6, 30))
        self.assertEqual(as_date("5 July 2025"), date(2025, 7, 5))
        self.assertIsNone(as_date(""))
        self.assertEqual(as_purchase_type("consignment"), "consignment")
        self.assertEqual(as_purchase_type("Purchased"), "purchased")
        self.assertEqual(as_purchase_type(""), "purchased")


class ConsignmentDueTests(TestCase):
    def setUp(self):
        from datetime import date, timedelta

        from apps.accounts.models import Employee
        from apps.catalogue.consignment import NEAR_DAYS
        from apps.catalogue.models import Category, Currency, ProductMaster, PurchaseType

        self.today = date(2026, 9, 11)
        cat = Category.objects.create(name="Jewellery", code="JW")
        cur = Currency.objects.create(code="USD", symbol="$")
        self.overdue = ProductMaster.objects.create(
            name="ANICH1",
            reference_id="ANICH-OVER",
            category=cat,
            currency=cur,
            product_type=PurchaseType.CONSIGNMENT,
            purchase_date=self.today - timedelta(days=61),
            due_date=self.today - timedelta(days=2),
        )
        self.soon = ProductMaster.objects.create(
            name="ANICH1",
            reference_id="ANICH-SOON",
            category=cat,
            currency=cur,
            product_type=PurchaseType.CONSIGNMENT,
            purchase_date=self.today,
            due_date=self.today + timedelta(days=NEAR_DAYS),
        )
        ProductMaster.objects.create(
            name="LATER",
            reference_id="ANICH-LATER",
            category=cat,
            currency=cur,
            product_type=PurchaseType.CONSIGNMENT,
            due_date=self.today + timedelta(days=NEAR_DAYS + 10),
        )
        ProductMaster.objects.create(
            name="BOUGHT",
            reference_id="BUY-1",
            category=cat,
            currency=cur,
            product_type=PurchaseType.PURCHASED,
            due_date=self.today - timedelta(days=1),
        )
        self.user = Employee.objects.create_user(
            employee_code="1001", password="test-pass", first_name="Test"
        )

    def test_summary_counts_overdue_and_near_not_purchased(self):
        from apps.catalogue.consignment import consignment_due_summary

        summary = consignment_due_summary(today=self.today)
        self.assertEqual(summary["overdue_count"], 1)
        self.assertEqual(summary["soon_count"], 1)
        self.assertEqual(summary["open_count"], 2)
        refs = {p.reference_id for p in summary["rows"]}
        self.assertEqual(refs, {"ANICH-OVER", "ANICH-SOON"})
        self.assertTrue(summary["rows"][0].is_overdue)

    def test_home_and_bell_show_real_consignment_rows(self):
        from unittest.mock import patch

        from django.test import override_settings

        self.client.force_login(self.user)
        static = {
            "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        }
        with override_settings(DEBUG=True, STORAGES=static):
            with patch("apps.catalogue.consignment.timezone.localdate", return_value=self.today):
                home = self.client.get("/")
                listed = self.client.get("/products/?purchase=consignment&due=open")
                detail = self.client.get(f"/products/{self.overdue.pk}/")
        self.assertEqual(home.status_code, 200)
        self.assertContains(home, "ANICH-OVER")
        self.assertContains(home, "Consignment due")
        self.assertContains(home, "ANICH-OVER — overdue 2d")
        self.assertContains(home, "View consignment dues")
        self.assertContains(home, "2 open")
        self.assertNotContains(home, "INV-004821")
        self.assertContains(home, "consignAlarmModal")
        self.assertEqual(listed.status_code, 200)
        self.assertContains(listed, "ANICH-OVER")
        self.assertContains(listed, "ANICH-SOON")
        self.assertNotContains(listed, "ANICH-LATER")
        self.assertNotContains(listed, "BUY-1")
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "Consignment")
        self.assertContains(detail, "Overdue 2d")


class ProductPhotoFilterTests(TestCase):
    def setUp(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from apps.accounts.models import Employee
        from apps.catalogue.models import Category, Currency, ProductImage, ProductMaster
        from apps.inventory.models import ProductItem
        from apps.locations.models import Location, LocationType

        cat = Category.objects.create(name="Jewellery", code="JW")
        cur = Currency.objects.create(code="USD", symbol="$")
        loc = Location.objects.create(name="Head office", code="HO", location_type=LocationType.HEAD_OFFICE)
        self.with_photo = ProductMaster.objects.create(
            name="SHOT", reference_id="PHOTO-1", category=cat, currency=cur,
        )
        self.no_photo = ProductMaster.objects.create(
            name="BLANK", reference_id="NOPIC-1", category=cat, currency=cur,
        )
        ProductImage.objects.create(
            product=self.with_photo,
            image=SimpleUploadedFile("ring.jpg", b"\xff\xd8\xff", content_type="image/jpeg"),
        )
        ProductItem.objects.create(barcode="PJPHOTO1", product=self.with_photo, location=loc)
        ProductItem.objects.create(barcode="PJNOPIC1", product=self.no_photo, location=loc)
        self.user = Employee.objects.create_user(
            employee_code="1001", password="test-pass", first_name="Test"
        )

    def test_list_can_keep_only_designs_with_photos(self):
        from django.test import override_settings

        self.client.force_login(self.user)
        static = {
            "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        }
        with override_settings(DEBUG=True, STORAGES=static):
            listed = self.client.get("/products/?media=photos")
            none = self.client.get("/products/?media=none")
        self.assertEqual(listed.status_code, 200)
        self.assertContains(listed, "PHOTO-1")
        self.assertNotContains(listed, "NOPIC-1")
        self.assertContains(listed, "Has photos")
        self.assertContains(none, "NOPIC-1")
        self.assertNotContains(none, "PHOTO-1")

    def test_piece_page_shows_design_photos(self):
        from django.test import override_settings

        self.client.force_login(self.user)
        static = {
            "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
            "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
        }
        with override_settings(DEBUG=True, STORAGES=static):
            page = self.client.get("/products/item/PJPHOTO1/")
            blank = self.client.get("/products/item/PJNOPIC1/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "idPhotoMain")
        self.assertContains(page, "same shots as the design")
        self.assertEqual(blank.status_code, 200)
        self.assertNotContains(blank, "idPhotoMain")


class ProductPhotoUploadTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import Permission
        from django.core.files.uploadedfile import SimpleUploadedFile

        from apps.accounts.models import Employee
        from apps.catalogue.models import Category, Currency, ProductMaster
        from apps.inventory.models import ProductItem
        from apps.locations.models import Location, LocationType

        cat = Category.objects.create(name="Jewellery", code="JW")
        cur = Currency.objects.create(code="USD", symbol="$")
        loc = Location.objects.create(name="Head office", code="HO", location_type=LocationType.HEAD_OFFICE)
        self.product = ProductMaster.objects.create(
            name="RING", reference_id="SUP-RING-1", category=cat, currency=cur,
        )
        ProductItem.objects.create(barcode="PJ99001", product=self.product, location=loc)
        self.user = Employee.objects.create_user(
            employee_code="2001", password="test-pass", first_name="Photo"
        )
        perm = Permission.objects.get(codename="can_upload_photos", content_type__app_label="catalogue")
        self.user.user_permissions.add(perm)
        self.jpeg = SimpleUploadedFile("shot.jpg", b"\xff\xd8\xff\xd9", content_type="image/jpeg")
        self.named = SimpleUploadedFile(
            "PJ99001 studio front.JPG", b"\xff\xd8\xff\xd9", content_type="image/jpeg",
        )

    def _static(self):
        from django.test import override_settings

        return override_settings(
            DEBUG=True,
            STORAGES={
                "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
                "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
            },
        )

    def test_forbidden_without_permission(self):
        from apps.accounts.models import Employee

        other = Employee.objects.create_user(
            employee_code="2002", password="test-pass", first_name="NoPhoto"
        )
        self.client.force_login(other)
        page = self.client.get("/products/photos/")
        self.assertEqual(page.status_code, 403)

    def test_lookup_by_barcode_then_upload(self):
        from apps.catalogue.models import ProductImage

        self.client.force_login(self.user)
        with self._static():
            lookup = self.client.get("/products/photos/?code=PJ99001")
            self.assertEqual(lookup.status_code, 200)
            self.assertContains(lookup, "PJ99001")
            self.assertContains(lookup, "SUP-RING-1")
            self.assertContains(lookup, "Attach photos to PJ99001")

            posted = self.client.post(
                "/products/photos/",
                {"mode": "upload", "code": "PJ99001", "photos": self.jpeg},
            )
            self.assertEqual(posted.status_code, 302)
            self.assertEqual(ProductImage.objects.filter(product=self.product).count(), 1)
            img = ProductImage.objects.get(product=self.product)
            self.assertTrue(img.is_primary)
            self.assertEqual(img.source_filename, "shot.jpg")

    def test_bulk_matches_pj_in_filename(self):
        from apps.catalogue.models import ProductImage

        self.client.force_login(self.user)
        with self._static():
            posted = self.client.post(
                "/products/photos/",
                {"mode": "bulk", "photos": self.named},
            )
            self.assertEqual(posted.status_code, 200)
            self.assertEqual(ProductImage.objects.filter(product=self.product).count(), 1)

    def test_unknown_code_does_not_create_stock(self):
        from apps.catalogue.models import ProductImage, ProductMaster

        before = ProductMaster.objects.count()
        self.client.force_login(self.user)
        with self._static():
            page = self.client.get("/products/photos/?code=PJ99999")
            self.assertEqual(page.status_code, 200)
            self.assertContains(page, "No design found")
        self.assertEqual(ProductMaster.objects.count(), before)
        self.assertEqual(ProductImage.objects.count(), 0)

    def test_can_remove_one_photo_and_all(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from apps.catalogue.models import ProductImage

        self.client.force_login(self.user)
        with self._static():
            self.client.post(
                "/products/photos/",
                {
                    "mode": "upload",
                    "code": "PJ99001",
                    "photos": [
                        SimpleUploadedFile("a.jpg", b"\xff\xd8\xff\xd9", content_type="image/jpeg"),
                        SimpleUploadedFile("b.jpg", b"\xff\xd8\xff\xd9", content_type="image/jpeg"),
                    ],
                },
            )
            self.assertEqual(ProductImage.objects.filter(product=self.product).count(), 2)
            keep = ProductImage.objects.filter(product=self.product).order_by("id").first()
            gone = ProductImage.objects.filter(product=self.product).order_by("id").last()
            deleted = self.client.post(
                "/products/photos/",
                {"mode": "delete", "code": "PJ99001", "image_id": str(gone.pk)},
            )
            self.assertEqual(deleted.status_code, 302)
            self.assertEqual(ProductImage.objects.filter(product=self.product).count(), 1)
            self.assertTrue(ProductImage.objects.filter(pk=keep.pk).exists())

            cleared = self.client.post(
                "/products/photos/",
                {"mode": "delete_all", "code": "PJ99001"},
            )
            self.assertEqual(cleared.status_code, 302)
            self.assertEqual(ProductImage.objects.filter(product=self.product).count(), 0)

    def test_oversized_upload_is_reported(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings

        from apps.catalogue.models import ProductImage

        self.client.force_login(self.user)
        big = SimpleUploadedFile("huge.jpg", b"\xff\xd8" + (b"\x00" * 2000) + b"\xd9", content_type="image/jpeg")
        with self._static(), override_settings(PRODUCT_PHOTO_MAX_UPLOAD_BYTES=500):
            posted = self.client.post(
                "/products/photos/",
                {"mode": "upload", "code": "PJ99001", "photos": big},
                follow=True,
            )
        self.assertEqual(posted.status_code, 200)
        self.assertEqual(ProductImage.objects.filter(product=self.product).count(), 0)
        self.assertContains(posted, "MB limit")

