"""Catch the jewellery Excel that staff upload on iadmin Product Master.

The live button opens website_product_reference.aspx. That file is a
workbook whose real header is row 2 of the "Jewellery Excel" sheet.
Each data row is one piece: a supplier code and a PJ barcode.

Purity is stored as the text they typed (DFLT - 18K becomes 18K).
A blank metal name is not treated as Silver — metal is inferred from
the purity text (18K → Gold, PT900 → Platinum).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Max

from apps.catalogue.models import (
    Category,
    Currency,
    Metal,
    ProductIntakeBatch,
    ProductIntakeLine,
    ProductMaster,
    PurchaseType,
    Purity,
    Supplier,
    strip_dflt_prefix,
)
from apps.core.legacy_import import as_date, as_purchase_type
from apps.inventory.models import ProductItem, StockStatus
from apps.locations.models import Location, LocationType

HEADER_ALIASES = {
    "pjnumber": "pj",
    "pj_number": "pj",
    "barcode": "pj",
    "supplier_product_code": "reference",
    "supplier_nick_name": "owner",
    "supplier": "supplier",
    "category": "category",
    "sub_category": "subcategory",
    "metal_name": "metal_name",
    "metal_purity": "purity",
    "metal_weight": "weight",
    "actual_price": "actual_price",
    "original_currency": "currency",
    "rate": "rate",
    "product_payment_type": "payment_type",
    "markup": "markup",
    "markup_amount": "markup_amount",
    "jewellry_size": "size",
    "jewellery_size": "size",
    "diamond_weight": "diamond_weight",
    "purchase_type": "purchase_type",
    "purchase_date": "purchase_date",
    "due_date": "due_date",
}

# Same header names as iadmin JewelleryExcelDData.xlsx / Cylver's upload sheet.
TEMPLATE_HEADERS = (
    "supplier",
    "Supplier_product_code",
    "Supplier_nick_name",
    "category",
    "sub_category",
    "Metal_name",
    "Metal Purity",
    "metal_weight",
    "actual_price",
    "Original_Currency",
    "rate",
    "product_payment_type",
    "markup",
    "markup_amount",
    "jewellry_size",
    "diamond_weight",
    "purchase_type",
    "purchase_date",
    "due_date",
    "PJNUMBER",
)

TEMPLATE_FILENAME = "JewelleryExcelDData.xlsx"


def _norm_header(value) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("/", " ").replace("-", " ")
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return HEADER_ALIASES.get(text, text)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decimal(value):
    text = _text(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _weight_text(value) -> str:
    number = _decimal(value)
    if number is None:
        return ""
    text = format(number.normalize(), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def infer_metal_name(metal_name: str, purity: str) -> str:
    typed = strip_dflt_prefix(metal_name)
    if typed:
        return typed
    purity = strip_dflt_prefix(purity).upper()
    if not purity:
        return ""
    if purity.startswith("PT") or "PLATIN" in purity:
        return "Platinum"
    if "SILVER" in purity or purity in {"SL", "SV", "AG"}:
        return "Silver"
    if re.search(r"\d+\s*K", purity) or "GOLD" in purity:
        return "Gold"
    return ""


def selling_price(actual, payment_type: str, markup, markup_amount):
    if actual is None:
        return None
    kind = (payment_type or "").strip().lower()
    if kind == "markuptype1" and markup is not None:
        return (actual * (Decimal("1") - (markup / Decimal("100")))).quantize(Decimal("0.01"))
    if kind == "markupamounttype1" and markup_amount is not None:
        return (actual - markup_amount).quantize(Decimal("0.01"))
    return actual


@dataclass
class IntakeRow:
    row_number: int
    pj: str
    reference: str
    owner: str
    supplier: str
    subcategory: str
    metal_name: str
    purity: str
    weight: str
    actual_price: Decimal | None
    currency: str
    rate: Decimal | None
    payment_type: str
    markup: Decimal | None
    markup_amount: Decimal | None
    size: str
    diamond_weight: str
    purchase_type: str = ""
    purchase_date: date | None = None
    due_date: date | None = None


@dataclass
class ParseResult:
    rows: list[IntakeRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    sheet: str = ""


def parse_jewellery_workbook(uploaded) -> ParseResult:
    from openpyxl import load_workbook

    result = ParseResult()
    try:
        wb = load_workbook(uploaded, data_only=True)
    except Exception as exc:
        result.errors.append(f"Could not open that file as Excel. {exc}")
        return result

    sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() == "jewellery excel":
            sheet = wb[name]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]
    result.sheet = sheet.title

    grid = []
    for i, row in enumerate(sheet.iter_rows(max_row=4, values_only=True), start=1):
        grid.append(row)
        if i >= 4:
            break
    header_idx = None
    for i, row in enumerate(grid):
        names = {_norm_header(cell) for cell in row}
        if "pj" in names or "reference" in names:
            header_idx = i
            break
    if header_idx is None:
        result.errors.append(
            "This is not the jewellery upload. Expected a Jewellery Excel sheet "
            "with PJNUMBER and Supplier_product_code on row 2."
        )
        return result

    headers = [_norm_header(cell) for cell in grid[header_idx]]
    start_row = header_idx + 2  # openpyxl rows are 1-based; header_idx is 0-based in grid

    seen = set()
    for offset, row in enumerate(sheet.iter_rows(min_row=start_row, values_only=True), start=start_row):
        data = {}
        for i, cell in enumerate(row):
            if i < len(headers) and headers[i]:
                data[headers[i]] = cell
        pj = _text(data.get("pj")).upper()
        reference = _text(data.get("reference"))
        if not pj and not reference:
            continue
        if not pj or not reference:
            result.errors.append(f"Row {offset}: needs both a PJ number and a supplier product code.")
            continue
        if pj in seen:
            result.errors.append(f"Row {offset}: {pj} is repeated in this file.")
            continue
        seen.add(pj)
        result.rows.append(IntakeRow(
            row_number=offset,
            pj=pj[:100],
            reference=reference[:100],
            owner=_text(data.get("owner")) or _text(data.get("supplier")) or pj,
            supplier=_text(data.get("supplier")),
            subcategory=_text(data.get("subcategory")).upper()[:100],
            metal_name=_text(data.get("metal_name")),
            purity=_text(data.get("purity")),
            weight=_weight_text(data.get("weight")),
            actual_price=_decimal(data.get("actual_price")),
            currency=(_text(data.get("currency")) or "USD").upper()[:10],
            rate=_decimal(data.get("rate")),
            payment_type=_text(data.get("payment_type")),
            markup=_decimal(data.get("markup")),
            markup_amount=_decimal(data.get("markup_amount")),
            size=_text(data.get("size"))[:50],
            diamond_weight=_weight_text(data.get("diamond_weight"))[:40],
            purchase_type=_text(data.get("purchase_type")),
            purchase_date=as_date(data.get("purchase_date")),
            due_date=as_date(data.get("due_date")),
        ))

    if not result.rows and not result.errors:
        result.errors.append("No data rows under the header.")
    return result


def default_location():
    loc = Location.objects.filter(location_type=LocationType.HEAD_OFFICE, is_active=True).first()
    if loc:
        return loc
    return Location.objects.filter(is_active=True).order_by("name").first()


def _category():
    cat = Category.objects.filter(code__iexact="JW").first()
    if cat:
        return cat
    cat = Category.objects.filter(name__iexact="jewellery").first()
    if cat:
        return cat
    return Category.objects.create(name="Jewellery", code="JW")


def _currency(code: str) -> Currency:
    code = (code or "USD").upper()[:10] or "USD"
    obj, _ = Currency.objects.get_or_create(code=code, defaults={"symbol": code})
    return obj


def _supplier(name: str) -> Supplier | None:
    name = (name or "").strip()
    if not name:
        return None
    obj = Supplier.objects.filter(name__iexact=name).first()
    if obj:
        return obj
    return Supplier.objects.create(name=name[:150])


def _metal_and_purity(metal_name: str, purity_text: str):
    purity_name = strip_dflt_prefix(purity_text)
    metal_label = infer_metal_name(metal_name, purity_text)
    metal = None
    purity = None
    if metal_label:
        metal, _ = Metal.objects.get_or_create(name=metal_label[:100])
    if purity_name and metal:
        purity, _ = Purity.objects.get_or_create(metal=metal, name=purity_name[:50])
    return metal, purity


def _lot_fields(row: IntakeRow, *, creating: bool) -> dict:
    fields = {}
    if row.purchase_type:
        fields["product_type"] = as_purchase_type(row.purchase_type)
    elif creating:
        fields["product_type"] = PurchaseType.PURCHASED
    if row.purchase_date:
        fields["purchase_date"] = row.purchase_date
    if row.due_date:
        fields["due_date"] = row.due_date
    return fields


def apply_rows(rows: list[IntakeRow], location: Location) -> dict:
    created = 0
    updated = 0
    notes = []
    lines = []
    category = _category()

    with transaction.atomic():
        for row in rows:
            metal, purity = _metal_and_purity(row.metal_name, row.purity)
            supplier = _supplier(row.supplier)
            currency = _currency(row.currency)
            price = selling_price(row.actual_price, row.payment_type, row.markup, row.markup_amount)
            weight = _decimal(row.weight)
            note = ""
            if row.purity and not purity:
                note = f"{row.pj}: purity “{row.purity}” was not stored — no metal name to attach it to."
            elif not row.weight:
                note = f"{row.pj}: no metal weight on this row."
            item = ProductItem.objects.filter(barcode__iexact=row.pj).select_related("product").first()
            if item:
                product = item.product
                product.reference_id = row.reference
                product.name = row.owner[:200]
                product.subcategory = row.subcategory
                product.currency = currency
                product.supplier = supplier
                product.metal = metal
                product.purity = purity
                product.net_weight = weight
                product.gold_weight = row.weight
                product.diamond_weight = row.diamond_weight
                product.size = row.size
                product.purchase_price = row.actual_price
                product.selling_price = price
                if row.rate is not None:
                    product.convert_rate = row.rate
                for key, value in _lot_fields(row, creating=False).items():
                    setattr(product, key, value)
                product.save()
                updated += 1
                action = ProductIntakeLine.Action.UPDATED
            else:
                product = ProductMaster.objects.create(
                    reference_id=row.reference,
                    name=row.owner[:200],
                    category=category,
                    subcategory=row.subcategory,
                    currency=currency,
                    supplier=supplier,
                    metal=metal,
                    purity=purity,
                    net_weight=weight,
                    gold_weight=row.weight,
                    diamond_weight=row.diamond_weight,
                    size=row.size,
                    purchase_price=row.actual_price,
                    selling_price=price,
                    convert_rate=row.rate if row.rate is not None else Decimal("1"),
                    **_lot_fields(row, creating=True),
                )
                item = ProductItem.objects.create(
                    barcode=row.pj,
                    product=product,
                    location=location,
                    status=StockStatus.PENDING,
                )
                created += 1
                action = ProductIntakeLine.Action.CREATED
            if note:
                notes.append(note)
            lines.append({
                "pj": row.pj,
                "reference": row.reference,
                "action": action,
                "item_id": item.pk,
                "row_number": row.row_number,
                "message": note,
            })

    return {
        "created": created,
        "updated": updated,
        "notes": notes[:12],
        "note_count": len(notes),
        "lines": lines,
    }


def build_jewellery_template() -> bytes:
    """Workbook staff fill — same sheet name and row-2 header as iadmin."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Jewellery Excel"
    ws.append([
        "Jewellery Excel — Perfect Jewel import template. Header is row 2. "
        "Each data row is one PJ code. Leave Metal_name blank if the sheet does."
    ])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(TEMPLATE_HEADERS))
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.append(list(TEMPLATE_HEADERS))
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D6E3F0")
    hints = {
        "supplier": "YZC",
        "Supplier_product_code": "Z-R4442",
        "Supplier_nick_name": "YZC1",
        "category": "JW",
        "sub_category": "RN",
        "Metal_name": "Leave blank if the sheet does",
        "Metal Purity": "DFLT - 18K",
        "metal_weight": "1.58",
        "actual_price": "239.40",
        "Original_Currency": "USD",
        "rate": "63",
        "product_payment_type": "markuptype1",
        "markup": "10",
        "purchase_type": "consignment or Purchased",
        "purchase_date": "11-09-2026",
        "due_date": "11-11-2026",
        "PJNUMBER": "PJ25107",
    }
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        hint = hints.get(header)
        if hint:
            ws.cell(2, col).comment = Comment(f"Example: {hint}", "PJ ERP")
        ws.column_dimensions[get_column_letter(col)].width = max(16, len(header) + 2)

    notes = wb.create_sheet("How to fill")
    notes.append(["Column", "What to put"])
    notes["A1"].font = Font(bold=True)
    notes["B1"].font = Font(bold=True)
    for header in TEMPLATE_HEADERS:
        notes.append([header, hints.get(header, "")])
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 42
    notes.append([])
    notes.append(["A new PJNUMBER creates a design and one piece. A PJ already in the catalog is updated, not duplicated."])
    notes.append(["DFLT - 18K is stored as 18K Gold. A blank Metal_name is not saved as Silver."])
    notes.append(["Start typing on row 3 of Jewellery Excel. Do not move the header off row 2."])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def record_intake_batch(
    *,
    user,
    location: Location | None,
    filename: str,
    raw_bytes: bytes | None,
    parsed: ParseResult | None,
    report: dict | None,
    source: str = ProductIntakeBatch.Source.EXCEL,
) -> ProductIntakeBatch:
    created = (report or {}).get("created", 0)
    updated = (report or {}).get("updated", 0)
    errors = list((parsed.errors if parsed else []) or [])
    applied = created + updated
    if applied and errors:
        status = ProductIntakeBatch.Status.PARTIAL
    elif applied:
        status = ProductIntakeBatch.Status.SUCCESS
    else:
        status = ProductIntakeBatch.Status.FAILED

    sno = (ProductIntakeBatch.objects.aggregate(n=Max("serial_no"))["n"] or 0) + 1
    notes = (report or {}).get("notes") or []
    batch = ProductIntakeBatch(
        serial_no=sno,
        source=source,
        filename=(filename or "")[:255],
        sheet=((parsed.sheet if parsed else "") or "")[:100],
        uploaded_by=user if getattr(user, "is_authenticated", False) else None,
        location=location,
        created_count=created,
        updated_count=updated,
        error_count=len(errors),
        status=status,
        notes="\n".join(notes)[:2000],
    )
    if raw_bytes:
        name = filename or TEMPLATE_FILENAME
        batch.workbook.save(name, ContentFile(raw_bytes), save=False)
    batch.save()

    rows = []
    for err in errors:
        rows.append(ProductIntakeLine(
            batch=batch,
            action=ProductIntakeLine.Action.ERROR,
            message=err[:500],
        ))
    for line in (report or {}).get("lines") or []:
        rows.append(ProductIntakeLine(
            batch=batch,
            barcode=(line.get("pj") or "")[:100],
            reference=(line.get("reference") or "")[:100],
            action=line["action"],
            item_id=line.get("item_id"),
            row_number=line.get("row_number") or 0,
            message=(line.get("message") or "")[:500],
        ))
    if rows:
        ProductIntakeLine.objects.bulk_create(rows)
    return batch

